import re
from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader
from docx import Document

def extract_text_from_pdf(file_path: str) -> str:
    """Extract plain text from a PDF file using pypdf."""
    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text.strip())
    return "\n\n".join(text_parts)

def extract_text_from_docx(file_path: str) -> str:
    """Extract plain text from a DOCX file using python-docx."""
    doc = Document(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())
    # Also check tables in docx
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
    return "\n\n".join(text_parts)

def extract_text(file_path: str) -> str:
    """Unified text extractor dispatcher for supported document formats."""
    p = Path(file_path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Expected PDF or DOCX.")

def split_into_clauses(raw_text: str) -> List[Dict[str, str]]:
    """
    Split agreement text into individual clause units.
    Uses regex to recognize numbered sections/clauses (e.g. 'Clause 4.1', 'Section 2.3', '1. Delivery Volumes')
    with a robust paragraph fallback.
    """
    cleaned_text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    
    # Primary regex pattern matching formal contract headings:
    # 1. 'Clause 1.1: ...' or 'Section 1.1 - ...' or 'Article 2: ...'
    # 2. '1.1 Title' or '2. Title'
    heading_pattern = re.compile(
        r'(?m)^(?:(?P<prefix>Section|Clause|Article)\s+(?P<num>[0-9IVXLCDM\.]+)|(?P<num_alt>[0-9]+\.[0-9]+|[0-9]+\.))\s*[:\-\.]?\s*(?P<title>[^\n]+)',
        re.IGNORECASE
    )

    matches = list(heading_pattern.finditer(cleaned_text))
    clauses: List[Dict[str, str]] = []

    if len(matches) >= 3:
        for i, match in enumerate(matches):
            start_pos = match.start()
            end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(cleaned_text)
            clause_content = cleaned_text[start_pos:end_pos].strip()

            ref = match.group("num") or match.group("num_alt") or f"{i+1}"
            prefix = match.group("prefix") or "Clause"
            title = (match.group("title") or "").strip()
            
            ref_str = f"{prefix} {ref}".strip() if match.group("prefix") else f"Clause {ref}".strip()

            if len(clause_content) > 20:
                clauses.append({
                    "clause_ref": ref_str,
                    "title": title or ref_str,
                    "text": clause_content
                })

    # Fallback: if structured headings weren't matched, split by double-newlines
    if len(clauses) < 3:
        paragraphs = [p.strip() for p in cleaned_text.split("\n\n") if len(p.strip()) > 35]
        clauses = []
        for idx, para in enumerate(paragraphs, start=1):
            first_line = para.split("\n")[0][:60]
            clauses.append({
                "clause_ref": f"Clause {idx}",
                "title": first_line,
                "text": para
            })

    return clauses
