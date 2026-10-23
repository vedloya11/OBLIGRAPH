import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

CONTRACT_TITLE = "MASTER COMMERCIAL SUPPLY & MATERIALS PROCUREMENT AGREEMENT"

CLAUSES = [
    (
        "Clause 1.1: Minimum Monthly Delivery Volumes",
        "Supplier covenants and agrees to manufacture, package, and deliver a minimum volume of 5,000 metric tons of battery-grade lithium carbonate per calendar month to the Buyer's Rotterdam receiving terminal on or before the fifth (5th) calendar day of each month. In the event delivered volumes fall below 5,000 metric tons or arrival is delayed beyond the specified monthly cutoff date, Supplier shall expedite replacement tonnage via air or priority freight at Supplier's sole cost and credit two percent (2%) per day of delay against the current monthly billing invoice."
    ),
    (
        "Clause 1.2: Hazardous Materials Packaging & Labeling",
        "All materials dispatched hereunder must be packaged in hermetically sealed UN-approved steel drums adhering to GHS hazard labeling classifications prior to container port hand-off. Any shipment found with ruptured container seals, leaking drums, or missing GHS hazard placards shall be subject to immediate quarantine, with all municipal port quarantine and hazmat repackaging surcharges charged directly to Supplier."
    ),
    (
        "Clause 2.1: Material Quality and Chemical Purity Standard",
        "Supplier guarantees that all lithium carbonate delivered under this Agreement shall satisfy a chemical purity of not less than 98.5% Li2CO3, with iron content strictly below 0.001% and moisture content below 0.25%. If laboratory analysis reveals chemical purity below 98.5%, Buyer may immediately reject the entire batch, and Supplier shall dispatch a conforming replacement batch within seventy-two (72) hours of rejection notification, reimbursing Buyer for any production downtime damages."
    ),
    (
        "Clause 2.2: Certificate of Chemical Analysis (COA) Submission",
        "Concurrent with the dispatch of each shipment lot, Supplier shall furnish an independent Certified Laboratory Purity Assay Certificate (COA) signed by an accredited metallurgical inspection agency. Failure to transmit a valid COA prior to port vessel arrival shall hold up customs clearance and suspend payment timelines without penalty or interest to Buyer."
    ),
    (
        "Clause 3.1: 30-Day Advance Written Notice for Price Adjustments",
        "Neither party may adjust the contracted baseline unit price without providing at least thirty (30) calendar days prior written notice to the other party via formal registered transmission. Any invoice issued reflecting increased unit prices without the satisfied 30-day advance notice shall be deemed void and legally non-binding, with prior contract pricing remaining in full effect."
    ),
    (
        "Clause 3.2: Quarterly Price Index Adjustment Cap",
        "Any price adjustment proposed pursuant to Clause 3.1 shall not exceed five percent (5.0%) of the preceding quarter's baseline index. Price increases exceeding the 5.0% threshold shall require mutual written sign-off and an escalated executive committee review within ten (10) business days."
    ),
    (
        "Clause 4.1: Inspection and Dispute Notification Window",
        "Buyer shall inspect incoming material consignments and provide formal written notice of any defect, damage, or quantity discrepancy within seven (7) business days following physical receipt at the discharge terminal. Failure to lodge written notice within said 7-day period shall constitute irrevocable acceptance of the consignment as delivered."
    ),
    (
        "Clause 5.1: Payment Settlement and Invoicing Terms",
        "Buyer shall remit full payment of undisputed invoice amounts Net thirty (30) days from receipt of a verified electronic invoice and matching Bill of Lading. Any undisputed balance remaining unpaid past 30 days shall accrue interest at SOFR plus 3.5% per annum until settled."
    ),
    (
        "Clause 6.1: Annual Carbon Intensity Audit and Environmental Standards",
        "Supplier shall submit an annual third-party ISO 14064 greenhouse gas audit report on or before December 15th of each contract year, verifying that supply chain carbon intensity does not exceed 2.4 kg CO2e per kg product. If carbon intensity exceeds this cap, Supplier shall purchase verified Gold Standard carbon offsets to balance the excess within thirty (30) days."
    ),
    (
        "Clause 7.1: Comprehensive Commercial General Liability Insurance",
        "Supplier shall continuously maintain Commercial General Liability insurance with an aggregate coverage limit of not less than $10,000,000 USD naming Buyer as an additional insured. Supplier must submit an updated ACORD certificate of insurance at least fifteen (15) days prior to the annual policy renewal date."
    )
]

def generate_docx(dest_path: Path):
    """Generate DOCX format of demo agreement."""
    doc = Document()
    
    # Title
    title = doc.add_heading(CONTRACT_TITLE, level=0)
    title.runs[0].font.size = Pt(18)
    title.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    p_preamble = doc.add_paragraph(
        "THIS MASTER SUPPLY AGREEMENT (the \"Agreement\") is entered into between Apex Materials Corp (\"Supplier\") "
        "and Nexus Industrial Technologies B.V. (\"Buyer\"). The parties agree to the following binding operational covenants:"
    )
    p_preamble.paragraph_format.space_after = Pt(14)

    for heading, text in CLAUSES:
        h = doc.add_heading(heading, level=2)
        h.runs[0].font.size = Pt(13)
        h.runs[0].font.color.rgb = RGBColor(37, 99, 235)
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(10)

    # Signatures
    doc.add_heading("IN WITNESS WHEREOF, the parties hereto have executed this Agreement.", level=3)
    p_sig = doc.add_paragraph("APEX MATERIALS CORP (Supplier)            NEXUS INDUSTRIAL TECHNOLOGIES (Buyer)\nBy: Authorized Officer                          By: Managing Director")
    p_sig.paragraph_format.space_before = Pt(18)

    doc.save(str(dest_path))
    print(f"Generated DOCX demo contract at: {dest_path}")

def generate_pdf(dest_path: Path):
    """Generate PDF format of demo agreement."""
    doc = SimpleDocTemplate(
        str(dest_path),
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=19,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=12
    )

    heading_style = ParagraphStyle(
        'ClauseHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1d4ed8'),
        spaceBefore=10,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )

    story = [
        Paragraph(CONTRACT_TITLE, title_style),
        Paragraph("<b>Parties:</b> Apex Materials Corp (Supplier) &amp; Nexus Industrial Technologies B.V. (Buyer)", body_style),
        HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceBefore=6, spaceAfter=10)
    ]

    for heading, text in CLAUSES:
        story.append(Paragraph(heading, heading_style))
        story.append(Paragraph(text, body_style))
        story.append(Spacer(1, 4))

    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#cbd5e1'), spaceBefore=8, spaceAfter=10))
    story.append(Paragraph("<b>Executed by Authorized Representatives of Supplier and Buyer.</b>", body_style))

    doc.build(story)
    print(f"Generated PDF demo contract at: {dest_path}")

def main():
    seed_dir = Path(__file__).resolve().parent
    seed_dir.mkdir(parents=True, exist_ok=True)

    docx_path = seed_dir / "demo_contract.docx"
    pdf_path = seed_dir / "demo_contract.pdf"

    generate_docx(docx_path)
    generate_pdf(pdf_path)

if __name__ == "__main__":
    main()
